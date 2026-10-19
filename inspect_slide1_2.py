import pptx
from pptx import Presentation

prs = Presentation('s:/SY/SEM 3/MP/EduVerse_AI_Review_2_Presentation.pptx')
for idx in [0, 1]:
    slide = prs.slides[idx]
    print(f"\n==================== SLIDE {idx+1} ====================")
    for s_idx, s in enumerate(slide.shapes):
        txt = s.text.strip().replace('\n', ' // ') if s.has_text_frame else ''
        xml = s._element.xml
        has_blip = 'blip' in xml
        print(f"Shape #{s_idx} '{s.name}' ({s.shape_type}) pos=({s.left/914400:.2f}, {s.top/914400:.2f}, {s.width/914400:.2f}, {s.height/914400:.2f}) blip={has_blip}: {txt[:80]}")
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for g_idx, gs in enumerate(s.shapes):
                gtxt = gs.text.strip().replace('\n', ' // ') if gs.has_text_frame else ''
                g_blip = 'blip' in gs._element.xml
                print(f"    Sub-shape #{g_idx} '{gs.name}' ({gs.shape_type}) pos=({gs.left/914400:.2f}, {gs.top/914400:.2f}) blip={g_blip}: {gtxt[:60]}")
