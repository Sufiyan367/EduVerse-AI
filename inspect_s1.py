import pptx
from pptx import Presentation

prs = Presentation('s:/SY/SEM 3/MP/EduVerse_AI_Review_2_Presentation.pptx')
s1 = prs.slides[0]
for s in s1.shapes:
    print(f"Shape: {s.name}, pos: ({s.left/914400:.2f}, {s.top/914400:.2f}, {s.width/914400:.2f}, {s.height/914400:.2f})")
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        for sub in s.shapes:
            print(f"   Sub: {sub.name}, pos: ({sub.left/914400:.2f}, {sub.top/914400:.2f}, {sub.width/914400:.2f}, {sub.height/914400:.2f})")
