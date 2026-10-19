import pptx
from pptx import Presentation

prs = Presentation('s:/SY/SEM 3/MP/EduVerse_AI_Review_2_Presentation.pptx')

s1 = prs.slides[0]
print("--- Slide 1 shapes ---")
for s in s1.shapes:
    print(f"Name: {s.name}, Type: {s.shape_type}, Pos: ({s.left/914400:.2f}, {s.top/914400:.2f}, {s.width/914400:.2f}, {s.height/914400:.2f})")
    if s.has_text_frame:
        print(f"  Text: {repr(s.text)}")
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        for sub in s.shapes:
            print(f"  Sub: {sub.name}, Pos: ({sub.left/914400:.2f}, {sub.top/914400:.2f}), Type: {sub.shape_type}")
            if sub.has_text_frame:
                print(f"    SubText: {repr(sub.text)}")

s8 = prs.slides[7]
print("\n--- Slide 8 shapes ---")
for s in s8.shapes:
    print(f"Name: {s.name}, Type: {s.shape_type}, Pos: ({s.left/914400:.2f}, {s.top/914400:.2f}, {s.width/914400:.2f}, {s.height/914400:.2f})")
    if s.has_text_frame:
        print(f"  Text: {repr(s.text)}")
