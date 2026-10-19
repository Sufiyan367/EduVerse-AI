import pptx
from pptx import Presentation

prs = Presentation('s:/SY/SEM 3/MP/EduVerse_AI_Review_2_Presentation.pptx')

for i in range(len(prs.slides)):
    s = prs.slides[i]
    print(f"\n--- Slide {i+1} background & template shapes ---")
    for shape in s.shapes:
        name = shape.name
        left, top, w, h = shape.left/914400, shape.top/914400, shape.width/914400, shape.height/914400
        # If it covers full width or near bottom/top
        if w > 15 or top > 9.5 or top < 1.0:
            txt = shape.text.strip().replace('\n', ' ') if shape.has_text_frame else ''
            print(f"  {name:20} pos=({left:.2f}, {top:.2f}, {w:.2f}, {h:.2f}) text='{txt[:50]}'")
