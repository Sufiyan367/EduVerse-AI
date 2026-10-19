import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -------------------------------------------------------------
# COLOR PALETTE
# -------------------------------------------------------------
COLOR_PRIMARY_NAVY = RGBColor(10, 37, 64)       # #0A2540
COLOR_DEEP_TEAL    = RGBColor(0, 95, 115)       # #005F73 (ICEM Primary)
COLOR_ACCENT_TEAL  = RGBColor(10, 147, 150)     # #0A9396
COLOR_SOFT_CYAN    = RGBColor(148, 210, 189)    # #94D2BD
COLOR_LIGHT_BG     = RGBColor(246, 250, 251)    # Clean off-white surface
COLOR_WHITE        = RGBColor(255, 255, 255)
COLOR_CARD_BORDER  = RGBColor(205, 226, 230)    # Soft teal border
COLOR_TEXT_DARK    = RGBColor(15, 23, 42)       # Slate 900
COLOR_TEXT_MUTED   = RGBColor(71, 85, 105)      # Slate 600
COLOR_AMBER        = RGBColor(202, 103, 2)      # Warm highlight
COLOR_SUCCESS      = RGBColor(45, 106, 79)      # Emerald green

FONT_HEADING = "Segoe UI"
FONT_BODY    = "Segoe UI"

def clear_slide_content(slide, keep_names):
    """Remove shapes except those in keep_names."""
    to_remove = [s for s in slide.shapes if s.name not in keep_names]
    for s in to_remove:
        sp = s._element
        sp.getparent().remove(sp)

def add_header(slide, title_text, category_text):
    """Add a clean, modern header that never collides with the Indira logo."""
    # Category / breadcrumb pill
    cat_box = slide.shapes.add_textbox(Inches(2.55), Inches(0.55), Inches(12.0), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.name = FONT_HEADING
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_ACCENT_TEAL

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(2.55), Inches(0.85), Inches(14.0), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_PRIMARY_NAVY

def add_footer(slide):
    """Add consistent official footer."""
    tb = slide.shapes.add_textbox(Inches(14.0), Inches(10.60), Inches(5.0), Inches(0.30))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "EduVerse AI • Mini Project Review-2"
    p.alignment = PP_ALIGN.RIGHT
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.0):
    """Create a rounded card container."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()
    return shape

print("Helper definitions loaded.")
