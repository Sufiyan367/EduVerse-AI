import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -------------------------------------------------------------
# PALETTE & CONSTANTS
# -------------------------------------------------------------
COLOR_PRIMARY_NAVY = RGBColor(10, 37, 64)       # #0A2540
COLOR_DEEP_TEAL    = RGBColor(0, 95, 115)       # #005F73 (Official ICEM Teal)
COLOR_ACCENT_TEAL  = RGBColor(10, 147, 150)     # #0A9396
COLOR_SOFT_CYAN    = RGBColor(148, 210, 189)    # #94D2BD
COLOR_VERY_LIGHT_CYAN = RGBColor(230, 245, 245) # #E6F5F5
COLOR_LIGHT_SURF   = RGBColor(248, 251, 252)    # Crisp light surface
COLOR_WHITE        = RGBColor(255, 255, 255)
COLOR_CARD_BORDER  = RGBColor(210, 230, 234)    # Soft teal border
COLOR_BORDER_TEAL  = RGBColor(148, 210, 189)    # Slightly stronger accent border
COLOR_TEXT_DARK    = RGBColor(15, 23, 42)       # Slate 900
COLOR_TEXT_BODY    = RGBColor(51, 65, 85)       # Slate 700
COLOR_TEXT_MUTED   = RGBColor(100, 116, 139)    # Slate 500
COLOR_AMBER        = RGBColor(180, 83, 9)       # Warm amber
COLOR_SUCCESS      = RGBColor(45, 106, 79)      # Emerald green
COLOR_TAG_BG       = RGBColor(235, 244, 246)

FONT_MAIN = "Segoe UI"

def clear_shapes_except(slide, keep_names):
    """Safely remove shapes whose names are not in keep_names."""
    to_remove = [s for s in slide.shapes if s.name not in keep_names]
    for s in to_remove:
        sp = s._element
        sp.getparent().remove(sp)

def add_header(slide, title_text, category_text):
    """Add standardized header that does not collide with the Indira logo."""
    cat_box = slide.shapes.add_textbox(Inches(2.60), Inches(0.48), Inches(12.0), Inches(0.32))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.name = FONT_MAIN
    p_c.font.size = Pt(10.5)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_ACCENT_TEAL

    title_box = slide.shapes.add_textbox(Inches(2.60), Inches(0.78), Inches(14.0), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_MAIN
    p_t.font.size = Pt(25)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_PRIMARY_NAVY

def add_footer(slide):
    """Add standard footer."""
    tb = slide.shapes.add_textbox(Inches(14.0), Inches(10.60), Inches(5.0), Inches(0.30))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "EduVerse AI • Mini Project Review-2"
    p.alignment = PP_ALIGN.RIGHT
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.0):
    """Create a rounded card container."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()
    return shape

# =============================================================
# SLIDE 1 BUILDER
# =============================================================
def build_slide_1(slide, assets_dir):
    # Slide 1: Premium Cover
    clear_shapes_except(slide, [])
    
    # Background subtle light fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(20.0), Inches(11.25))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(250, 252, 253)
    bg.line.fill.background()
    
    # Official Indira Banner Logo at top
    banner_path = os.path.join(assets_dir, 'image8.png')
    if os.path.exists(banner_path):
        slide.shapes.add_picture(banner_path, Inches(1.2), Inches(0.65), Inches(5.8), Inches(1.8))
        
    # Department & Academic Year Header Pill
    pill = add_card(slide, Inches(1.2), Inches(2.75), Inches(8.4), Inches(0.48), bg_color=COLOR_DEEP_TEAL, border_color=None)
    tf_p = pill.text_frame
    tf_p.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_p.margin_left = Inches(0.2)
    p_p = tf_p.paragraphs[0]
    p_p.text = "MINI PROJECT REVIEW – 2  •  ACADEMIC YEAR 2026–2027"
    p_p.font.name = FONT_MAIN
    p_p.font.size = Pt(11)
    p_p.font.bold = True
    p_p.font.color.rgb = COLOR_WHITE
    
    # Project Title
    t_box = slide.shapes.add_textbox(Inches(1.2), Inches(3.40), Inches(8.8), Inches(1.4))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = "EduVerse AI"
    p_t.font.name = FONT_MAIN
    p_t.font.size = Pt(54)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_PRIMARY_NAVY
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.75), Inches(8.8), Inches(0.9))
    tf_s = sub_box.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
    p_s = tf_s.paragraphs[0]
    p_s.text = "An Animated Virtual Tutor for Personalized Learning"
    p_s.font.name = FONT_MAIN
    p_s.font.size = Pt(21)
    p_s.font.bold = True
    p_s.font.color.rgb = COLOR_DEEP_TEAL
    
    # Department details
    dept_box = slide.shapes.add_textbox(Inches(1.2), Inches(5.80), Inches(8.8), Inches(0.9))
    tf_d = dept_box.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = tf_d.margin_right = tf_d.margin_top = tf_d.margin_bottom = 0
    p_d1 = tf_d.paragraphs[0]
    p_d1.text = "Department of Artificial Intelligence & Data Science"
    p_d1.font.name = FONT_MAIN
    p_d1.font.size = Pt(15)
    p_d1.font.bold = True
    p_d1.font.color.rgb = COLOR_TEXT_DARK
    p_d2 = tf_d.add_paragraph()
    p_d2.text = "Indira College of Engineering & Management, Pune (Autonomous)"
    p_d2.font.name = FONT_MAIN
    p_d2.font.size = Pt(13)
    p_d2.font.color.rgb = COLOR_TEXT_MUTED
    p_d2.space_before = Pt(4)
    
    # Presenters Card
    pres_card = add_card(slide, Inches(1.2), Inches(7.10), Inches(4.3), Inches(2.70), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf_pr = pres_card.text_frame
    tf_pr.margin_left = Inches(0.3)
    tf_pr.margin_top = Inches(0.25)
    p_pr_h = tf_pr.paragraphs[0]
    p_pr_h.text = "PRESENTED BY"
    p_pr_h.font.name = FONT_MAIN
    p_pr_h.font.size = Pt(11)
    p_pr_h.font.bold = True
    p_pr_h.font.color.rgb = COLOR_ACCENT_TEAL
    
    members = ["Aryan Shinde", "Huzaifa Sheikh", "Sufiyan Surve"]
    for m in members:
        p_m = tf_pr.add_paragraph()
        p_m.text = f"•  {m}"
        p_m.font.name = FONT_MAIN
        p_m.font.size = Pt(14)
        p_m.font.bold = True
        p_m.font.color.rgb = COLOR_PRIMARY_NAVY
        p_m.space_before = Pt(6)
        
    p_sub = tf_pr.add_paragraph()
    p_sub.text = "B.E. Artificial Intelligence & Data Science"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(10.5)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_sub.space_before = Pt(8)
    
    # Mentor Card
    men_card = add_card(slide, Inches(5.75), Inches(7.10), Inches(4.25), Inches(2.70), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf_me = men_card.text_frame
    tf_me.margin_left = Inches(0.3)
    tf_me.margin_top = Inches(0.25)
    p_me_h = tf_me.paragraphs[0]
    p_me_h.text = "UNDER THE GUIDANCE OF"
    p_me_h.font.name = FONT_MAIN
    p_me_h.font.size = Pt(11)
    p_me_h.font.bold = True
    p_me_h.font.color.rgb = COLOR_ACCENT_TEAL
    
    p_me_n = tf_me.add_paragraph()
    p_me_n.text = "Prof. Monika Patil"
    p_me_n.font.name = FONT_MAIN
    p_me_n.font.size = Pt(16)
    p_me_n.font.bold = True
    p_me_n.font.color.rgb = COLOR_PRIMARY_NAVY
    p_me_n.space_before = Pt(8)
    
    p_me_d = tf_me.add_paragraph()
    p_me_d.text = "Assistant Professor & Project Mentor\nDept. of Artificial Intelligence & Data Science\nICEM, Pune"
    p_me_d.font.name = FONT_MAIN
    p_me_d.font.size = Pt(11.5)
    p_me_d.font.color.rgb = COLOR_TEXT_BODY
    p_me_d.space_before = Pt(6)

    # Right Hero Container
    hero_card = add_card(slide, Inches(10.35), Inches(0.85), Inches(8.65), Inches(9.20), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.5)
    hero_path = os.path.join(assets_dir, 'hero_visual.jpg')
    if os.path.exists(hero_path):
        slide.shapes.add_picture(hero_path, Inches(10.55), Inches(1.05), Inches(8.25), Inches(8.20))
        
    # Hero Caption Tag
    cap_tag = add_card(slide, Inches(11.2), Inches(9.38), Inches(6.9), Inches(0.50), bg_color=COLOR_PRIMARY_NAVY, border_color=None)
    tf_cp = cap_tag.text_frame
    tf_cp.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_cp = tf_cp.paragraphs[0]
    p_cp.text = "AI-Assisted Contextual Learning & Virtual Tutoring System"
    p_cp.alignment = PP_ALIGN.CENTER
    p_cp.font.name = FONT_MAIN
    p_cp.font.size = Pt(11)
    p_cp.font.bold = True
    p_cp.font.color.rgb = COLOR_WHITE


# =============================================================
# SLIDE 2 BUILDER
# =============================================================
def build_slide_2(slide):
    # Slide 2: Agenda
    clear_shapes_except(slide, ['Freeform 2', 'Group 42'])
    add_header(slide, "Agenda", "Project Review – 2 • Presentation Structure")
    add_footer(slide)
    
    agenda_items = [
        ("01", "Introduction", "Problem context, fragmented study resources, and proposed virtual tutor concept.", "Idea & Problem"),
        ("02", "Aim & Objectives", "Primary project goal and 6 core functional engineering deliverables.", "Target & Goals"),
        ("03", "Tools Used", "Proposed technology stack across UI, API, LLM, document engine, and 3D avatar.", "Tech Stack"),
        ("04", "Proposed Solution / Overview", "End-to-end system architecture, PDF ingestion pipeline, and data flow.", "System Flow"),
        ("05", "Implementation / Prototype", "Interactive web UI demonstration, dialogue stream, and active quiz module.", "Functional UI"),
        ("06", "Applications", "Practical student self-learning, exam revision, and interactive practice use cases.", "Practical Value"),
        ("07", "Future Scope", "Phased development roadmap, voice integration, and adaptive difficulty.", "Roadmap"),
        ("08", "Conclusion & References", "Project summary, key takeaways, and verified technical documentation.", "Wrap-up & Docs")
    ]
    
    card_w = Inches(4.25)
    card_h = Inches(3.70)
    gap_x = Inches(0.25)
    start_x = Inches(1.0)
    
    for i, (num, title, desc, tag) in enumerate(agenda_items):
        row = i // 4
        col = i % 4
        left = start_x + col * (card_w + gap_x)
        top = Inches(1.85) if row == 0 else Inches(5.85)
        
        card = add_card(slide, left, top, card_w, card_h, bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.2)
        tf = card.text_frame
        tf.margin_left = Inches(0.28)
        tf.margin_right = Inches(0.28)
        tf.margin_top = Inches(0.26)
        tf.word_wrap = True
        
        # Header with number pill
        p_num = tf.paragraphs[0]
        p_num.text = f"SECTION  {num}"
        p_num.font.name = FONT_MAIN
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT_TEAL
        
        # Section Title
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_MAIN
        p_t.font.size = Pt(16.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY_NAVY
        p_t.space_before = Pt(4)
        
        # Tag pill
        p_tag = tf.add_paragraph()
        p_tag.text = f"FOCUS: {tag.upper()}"
        p_tag.font.name = FONT_MAIN
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_DEEP_TEAL
        p_tag.space_before = Pt(4)
        
        # Description
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_MAIN
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_BODY
        p_d.space_before = Pt(8)


# =============================================================
# SLIDE 3 BUILDER
# =============================================================
def build_slide_3(slide):
    # Slide 3: Introduction
    clear_shapes_except(slide, ['Freeform 2', 'Group 3'])
    add_header(slide, "Introduction", "Context & Problem Statement")
    add_footer(slide)
    
    # Column 1: THE PROBLEM
    col1 = add_card(slide, Inches(1.0), Inches(1.75), Inches(5.6), Inches(6.5), bg_color=COLOR_WHITE, border_color=RGBColor(248, 180, 180), border_width=1.5)
    tf1 = col1.text_frame
    tf1.margin_left = tf1.margin_right = Inches(0.30)
    tf1.margin_top = Inches(0.28)
    tf1.word_wrap = True
    
    p1_badge = tf1.paragraphs[0]
    p1_badge.text = "THE PROBLEM"
    p1_badge.font.name = FONT_MAIN
    p1_badge.font.size = Pt(11)
    p1_badge.font.bold = True
    p1_badge.font.color.rgb = RGBColor(185, 28, 28)
    
    p1_h = tf1.add_paragraph()
    p1_h.text = "Limitations of Static Learning"
    p1_h.font.name = FONT_MAIN
    p1_h.font.size = Pt(18)
    p1_h.font.bold = True
    p1_h.font.color.rgb = COLOR_PRIMARY_NAVY
    p1_h.space_before = Pt(2)
    
    pts1 = [
        ("Fragmented Study Material", "Students depend on disconnected textbook scans, static PDFs, and lengthy video lectures without unified assistance."),
        ("Difficult Contextual Doubts", "Finding clear explanations for specific course doubts requires switching across generic web platforms, losing syllabus focus."),
        ("No Instant Practice Loop", "Passive reading material does not automatically convert into interactive exercises, flashcards, or practice questions.")
    ]
    for title, desc in pts1:
        p_pt = tf1.add_paragraph()
        p_pt.text = f"•  {title}"
        p_pt.font.name = FONT_MAIN
        p_pt.font.size = Pt(13)
        p_pt.font.bold = True
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_before = Pt(12)
        
        p_pd = tf1.add_paragraph()
        p_pd.text = desc
        p_pd.font.name = FONT_MAIN
        p_pd.font.size = Pt(11.5)
        p_pd.font.color.rgb = COLOR_TEXT_BODY
        p_pd.space_before = Pt(2)
        
    # Column 2: THE GAP
    col2 = add_card(slide, Inches(6.85), Inches(1.75), Inches(5.6), Inches(6.5), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.5)
    tf2 = col2.text_frame
    tf2.margin_left = tf2.margin_right = Inches(0.30)
    tf2.margin_top = Inches(0.28)
    tf2.word_wrap = True
    
    p2_badge = tf2.paragraphs[0]
    p2_badge.text = "THE GAP"
    p2_badge.font.name = FONT_MAIN
    p2_badge.font.size = Pt(11)
    p2_badge.font.bold = True
    p2_badge.font.color.rgb = COLOR_AMBER
    
    p2_h = tf2.add_paragraph()
    p2_h.text = "The Learning Bottleneck"
    p2_h.font.name = FONT_MAIN
    p2_h.font.size = Pt(18)
    p2_h.font.bold = True
    p2_h.font.color.rgb = COLOR_PRIMARY_NAVY
    p2_h.space_before = Pt(2)
    
    # 3 mini cards inside column 2
    gap_boxes = [
        ("📄 Static Content", "Unstructured lecture slides, dense textbooks, and static PDF notes.", Inches(3.0)),
        ("⏸️ Limited Interaction", "One-way passive reading with zero real-time question answering.", Inches(4.5)),
        ("⚠️ Generic Explanations", "Non-personalized explanations without student-paced breakdown.", Inches(6.0))
    ]
    for b_title, b_desc, b_top in gap_boxes:
        gb = add_card(slide, Inches(7.15), b_top, Inches(5.0), Inches(1.15), bg_color=COLOR_TAG_BG, border_color=COLOR_CARD_BORDER, border_width=1.0)
        tf_gb = gb.text_frame
        tf_gb.margin_left = tf_gb.margin_right = Inches(0.18)
        tf_gb.margin_top = Inches(0.12)
        p_gb1 = tf_gb.paragraphs[0]
        p_gb1.text = b_title
        p_gb1.font.name = FONT_MAIN
        p_gb1.font.size = Pt(12)
        p_gb1.font.bold = True
        p_gb1.font.color.rgb = COLOR_PRIMARY_NAVY
        p_gb2 = tf_gb.add_paragraph()
        p_gb2.text = b_desc
        p_gb2.font.name = FONT_MAIN
        p_gb2.font.size = Pt(10.5)
        p_gb2.font.color.rgb = COLOR_TEXT_BODY
        p_gb2.space_before = Pt(2)
        
    # Downward arrows between gap boxes
    a1 = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.5), Inches(4.20), Inches(0.3), Inches(0.24))
    a1.fill.solid()
    a1.fill.fore_color.rgb = COLOR_ACCENT_TEAL
    a1.line.fill.background()
    a2 = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.5), Inches(5.70), Inches(0.3), Inches(0.24))
    a2.fill.solid()
    a2.fill.fore_color.rgb = COLOR_ACCENT_TEAL
    a2.line.fill.background()

    # Column 3: OUR IDEA
    col3 = add_card(slide, Inches(12.70), Inches(1.75), Inches(6.3), Inches(6.5), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf3 = col3.text_frame
    tf3.margin_left = tf3.margin_right = Inches(0.30)
    tf3.margin_top = Inches(0.28)
    tf3.word_wrap = True
    
    p3_badge = tf3.paragraphs[0]
    p3_badge.text = "OUR IDEA"
    p3_badge.font.name = FONT_MAIN
    p3_badge.font.size = Pt(11)
    p3_badge.font.bold = True
    p3_badge.font.color.rgb = COLOR_DEEP_TEAL
    
    p3_h = tf3.add_paragraph()
    p3_h.text = "EduVerse AI Virtual Tutor"
    p3_h.font.name = FONT_MAIN
    p3_h.font.size = Pt(18)
    p3_h.font.bold = True
    p3_h.font.color.rgb = COLOR_PRIMARY_NAVY
    p3_h.space_before = Pt(2)
    
    p3_sub = tf3.add_paragraph()
    p3_sub.text = "An interactive AI tutor bridging course documents with animated virtual guidance:"
    p3_sub.font.name = FONT_MAIN
    p3_sub.font.size = Pt(11.5)
    p3_sub.font.color.rgb = COLOR_TEXT_BODY
    p3_sub.space_before = Pt(6)
    
    pillars = [
        ("📚 Study Material", "Direct upload and ingestion of syllabus notes and textbook chapters."),
        ("🧠 AI Understanding", "Contextual parsing of text to anchor responses in student documents."),
        ("💬 Interactive Q&A", "Real-time conversational explanations tailored to student doubts."),
        ("📝 Quiz Generation", "Automated MCQ creation from study content for instant practice."),
        ("🤖 Animated Tutor", "A WebGL 3D virtual avatar interface that increases engagement.")
    ]
    for pil_t, pil_d in pillars:
        p_pt = tf3.add_paragraph()
        p_pt.text = f"{pil_t}: {pil_d}"
        p_pt.font.name = FONT_MAIN
        p_pt.font.size = Pt(11)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_before = Pt(5)

    # Bottom Full-Width Workflow Ribbon
    ribbon = add_card(slide, Inches(1.0), Inches(8.50), Inches(18.0), Inches(1.70), bg_color=COLOR_LIGHT_SURF, border_color=COLOR_BORDER_TEAL, border_width=1.2)
    tf_r = ribbon.text_frame
    tf_r.margin_left = tf_r.margin_right = Inches(0.3)
    tf_r.margin_top = Inches(0.18)
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "CORE INTERACTIVE LEARNING PIPELINE"
    p_rh.font.name = FONT_MAIN
    p_rh.font.size = Pt(10.5)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_ACCENT_TEAL
    
    # 6 steps
    steps = ["1. UPLOAD PDF", "2. UNDERSTAND", "3. ASK DOUBT", "4. EXPLAIN", "5. PRACTICE QUIZ", "6. REVIEW FEEDBACK"]
    step_w = Inches(2.65)
    step_gap = Inches(0.24)
    for s_idx, st in enumerate(steps):
        s_left = Inches(1.3) + s_idx * (step_w + step_gap)
        sc = add_card(slide, s_left, Inches(9.10), step_w, Inches(0.85), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.0)
        tf_sc = sc.text_frame
        tf_sc.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_sc = tf_sc.paragraphs[0]
        p_sc.text = st
        p_sc.alignment = PP_ALIGN.CENTER
        p_sc.font.name = FONT_MAIN
        p_sc.font.size = Pt(11.5)
        p_sc.font.bold = True
        p_sc.font.color.rgb = COLOR_PRIMARY_NAVY
        
        # small arrow between cards
        if s_idx < len(steps) - 1:
            arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, s_left + step_w + Inches(0.04), Inches(9.40), Inches(0.16), Inches(0.22))
            arr.fill.solid()
            arr.fill.fore_color.rgb = COLOR_ACCENT_TEAL
            arr.line.fill.background()


# =============================================================
# SLIDE 4 BUILDER
# =============================================================
def build_slide_4(slide):
    # Slide 4: Aim & Objectives
    clear_shapes_except(slide, ['Freeform 2', 'Group 3'])
    add_header(slide, "Aim & Objectives", "Project Goal & Engineering Deliverables")
    add_footer(slide)
    
    # Top AIM Banner
    aim_card = add_card(slide, Inches(1.0), Inches(1.70), Inches(18.0), Inches(1.90), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf_a = aim_card.text_frame
    tf_a.margin_left = tf_a.margin_right = Inches(0.35)
    tf_a.margin_top = Inches(0.22)
    tf_a.word_wrap = True
    
    p_ab = tf_a.paragraphs[0]
    p_ab.text = "CORE PROJECT AIM"
    p_ab.font.name = FONT_MAIN
    p_ab.font.size = Pt(11)
    p_ab.font.bold = True
    p_ab.font.color.rgb = COLOR_ACCENT_TEAL
    
    p_at = tf_a.add_paragraph()
    p_at.text = "“To develop an AI-powered animated virtual tutor that helps students interact with their study material, obtain contextual explanations and practice concepts through automatically generated quizzes.”"
    p_at.font.name = FONT_MAIN
    p_at.font.size = Pt(17.5)
    p_at.font.bold = True
    p_at.font.color.rgb = COLOR_PRIMARY_NAVY
    p_at.space_before = Pt(4)
    
    # 6 Objective Cards (2 rows x 3 columns)
    objs = [
        ("01", "Interactive AI Tutor", "Provide an engaging virtual learning assistant capable of answering student queries in a natural, conversational manner.", "Conversation"),
        ("02", "Study Material Understanding", "Parse and extract machine-readable text from student-uploaded PDF documents for contextual grounding.", "Document Parsing"),
        ("03", "Contextual Q&A", "Deliver accurate, document-relevant explanations tailored to specific student doubts, minimizing generic answers.", "Targeted Doubts"),
        ("04", "Simplified Explanations", "Deconstruct complex theoretical definitions and formulas into structured, student-friendly conceptual analogies.", "Clarity & Depth"),
        ("05", "AI Quiz Generation", "Automatically generate multiple-choice practice questions directly from the uploaded material for self-testing.", "Active Assessment"),
        ("06", "Learning Feedback", "Provide instant scoring, answer rationales, and self-assessment insights to reinforce student understanding.", "Evaluation")
    ]
    
    card_w = Inches(5.80)
    card_h = Inches(3.10)
    start_x = Inches(1.0)
    gap_x = Inches(0.30)
    
    for i, (num, title, desc, tag) in enumerate(objs):
        row = i // 3
        col = i % 3
        left = start_x + col * (card_w + gap_x)
        top = Inches(3.85) if row == 0 else Inches(7.15)
        
        card = add_card(slide, left, top, card_w, card_h, bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.2)
        tf = card.text_frame
        tf.margin_left = tf.margin_right = Inches(0.28)
        tf.margin_top = Inches(0.22)
        tf.word_wrap = True
        
        # Pill badge
        p_n = tf.paragraphs[0]
        p_n.text = f"OBJECTIVE  {num}  •  {tag.upper()}"
        p_n.font.name = FONT_MAIN
        p_n.font.size = Pt(10)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_ACCENT_TEAL
        
        # Title
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_MAIN
        p_t.font.size = Pt(16.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY_NAVY
        p_t.space_before = Pt(3)
        
        # Description
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_MAIN
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_BODY
        p_d.space_before = Pt(8)


# =============================================================
# SLIDE 5 BUILDER
# =============================================================
def build_slide_5(slide):
    # Slide 5: System Architecture & Technology Stack
    clear_shapes_except(slide, ['Freeform 2', 'Group 3'])
    add_header(slide, "System Architecture & Technology Stack", "Engineering Design & End-to-End Pipeline")
    add_footer(slide)
    
    # Left Column: PROPOSED TECHNOLOGY STACK (~30% width)
    stack_card = add_card(slide, Inches(1.0), Inches(1.70), Inches(5.3), Inches(8.5), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf_s = stack_card.text_frame
    tf_s.margin_left = tf_s.margin_right = Inches(0.26)
    tf_s.margin_top = Inches(0.24)
    tf_s.word_wrap = True
    
    p_sb = tf_s.paragraphs[0]
    p_sb.text = "PROPOSED TECHNOLOGY STACK"
    p_sb.font.name = FONT_MAIN
    p_sb.font.size = Pt(11)
    p_sb.font.bold = True
    p_sb.font.color.rgb = COLOR_ACCENT_TEAL
    
    techs = [
        ("Frontend Interface", "React", "Component-based single page application & interactive dashboard."),
        ("Backend Services", "Python / FastAPI", "High-performance asynchronous REST API routing and logic."),
        ("AI Processing Layer", "LLM API", "Contextual reasoning, explanation generation & quiz creation."),
        ("Document Processing", "PyMuPDF", "High-speed text extraction from uploaded student PDF files."),
        ("Database Layer", "SQLite / PostgreSQL", "Session state, document metadata, and quiz scores storage."),
        ("Avatar Technology", "Three.js / WebGL", "Web-based animated 3D virtual assistant display."),
        ("VCS & API Testing", "Git, GitHub & Postman", "Version control workflow and automated API verification.")
    ]
    
    for layer, tech, role in techs:
        p_l = tf_s.add_paragraph()
        p_l.text = f"{layer}:  {tech}"
        p_l.font.name = FONT_MAIN
        p_l.font.size = Pt(11.5)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_PRIMARY_NAVY
        p_l.space_before = Pt(8)
        
        p_r = tf_s.add_paragraph()
        p_r.text = role
        p_r.font.name = FONT_MAIN
        p_r.font.size = Pt(10)
        p_r.font.color.rgb = COLOR_TEXT_MUTED
        p_r.space_before = Pt(1)

    # Right Canvas: SYSTEM ARCHITECTURE (~70% width)
    arch_card = add_card(slide, Inches(6.60), Inches(1.70), Inches(12.4), Inches(8.5), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.5)
    tf_a = arch_card.text_frame
    tf_a.margin_left = Inches(0.35)
    tf_a.margin_top = Inches(0.22)
    p_ab = tf_a.paragraphs[0]
    p_ab.text = "END-TO-END SYSTEM ARCHITECTURE"
    p_ab.font.name = FONT_MAIN
    p_ab.font.size = Pt(11)
    p_ab.font.bold = True
    p_ab.font.color.rgb = COLOR_DEEP_TEAL
    
    # Architecture Flow Boxes
    # Tier 1: Client / User
    tier1_user = add_card(slide, Inches(7.0), Inches(2.35), Inches(3.4), Inches(1.15), bg_color=COLOR_TAG_BG, border_color=COLOR_CARD_BORDER)
    tf_u = tier1_user.text_frame
    tf_u.margin_left = tf_u.margin_right = Inches(0.15)
    tf_u.margin_top = Inches(0.12)
    p_u1 = tf_u.paragraphs[0]
    p_u1.text = "👨‍🎓 Student User"
    p_u1.font.name = FONT_MAIN
    p_u1.font.size = Pt(12)
    p_u1.font.bold = True
    p_u1.font.color.rgb = COLOR_PRIMARY_NAVY
    p_u2 = tf_u.add_paragraph()
    p_u2.text = "Uploads PDF • Asks Doubts • Takes Quizzes"
    p_u2.font.name = FONT_MAIN
    p_u2.font.size = Pt(10)
    p_u2.font.color.rgb = COLOR_TEXT_BODY
    
    # Tier 1 to 2 arrow
    arr1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(10.55), Inches(2.78), Inches(0.4), Inches(0.24))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = COLOR_ACCENT_TEAL
    arr1.line.fill.background()
    
    # Tier 1: Frontend
    tier1_ui = add_card(slide, Inches(11.1), Inches(2.35), Inches(3.6), Inches(1.15), bg_color=COLOR_TAG_BG, border_color=COLOR_CARD_BORDER)
    tf_ui = tier1_ui.text_frame
    tf_ui.margin_left = tf_ui.margin_right = Inches(0.15)
    tf_ui.margin_top = Inches(0.12)
    p_ui1 = tf_ui.paragraphs[0]
    p_ui1.text = "🖥️ React Web Interface"
    p_ui1.font.name = FONT_MAIN
    p_ui1.font.size = Pt(12)
    p_ui1.font.bold = True
    p_ui1.font.color.rgb = COLOR_PRIMARY_NAVY
    p_ui2 = tf_ui.add_paragraph()
    p_ui2.text = "Dashboard • Dialogue Stream • Avatar Canvas"
    p_ui2.font.name = FONT_MAIN
    p_ui2.font.size = Pt(10)
    p_ui2.font.color.rgb = COLOR_TEXT_BODY

    # Tier 1 to 2 arrow (to Backend)
    arr2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(14.85), Inches(2.78), Inches(0.4), Inches(0.24))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = COLOR_ACCENT_TEAL
    arr2.line.fill.background()
    
    # Tier 2: Backend API
    tier2_api = add_card(slide, Inches(15.4), Inches(2.35), Inches(3.2), Inches(1.15), bg_color=COLOR_VERY_LIGHT_CYAN, border_color=COLOR_BORDER_TEAL)
    tf_api = tier2_api.text_frame
    tf_api.margin_left = tf_api.margin_right = Inches(0.15)
    tf_api.margin_top = Inches(0.12)
    p_ap1 = tf_api.paragraphs[0]
    p_ap1.text = "⚙️ FastAPI Backend"
    p_ap1.font.name = FONT_MAIN
    p_ap1.font.size = Pt(12)
    p_ap1.font.bold = True
    p_ap1.font.color.rgb = COLOR_DEEP_TEAL
    p_ap2 = tf_api.add_paragraph()
    p_ap2.text = "REST Endpoints • Session Coordinator"
    p_ap2.font.name = FONT_MAIN
    p_ap2.font.size = Pt(10)
    p_ap2.font.color.rgb = COLOR_TEXT_BODY

    # Middle Level: Side Ingestion & Core AI
    # Side Ingestion Card
    side_doc = add_card(slide, Inches(7.0), Inches(4.0), Inches(5.4), Inches(1.50), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER)
    tf_sd = side_doc.text_frame
    tf_sd.margin_left = Inches(0.2)
    tf_sd.margin_top = Inches(0.15)
    p_sd1 = tf_sd.paragraphs[0]
    p_sd1.text = "📑 Document Processing Pipeline (PyMuPDF)"
    p_sd1.font.name = FONT_MAIN
    p_sd1.font.size = Pt(11.5)
    p_sd1.font.bold = True
    p_sd1.font.color.rgb = COLOR_PRIMARY_NAVY
    p_sd2 = tf_sd.add_paragraph()
    p_sd2.text = "1. Student uploads syllabus PDF / lecture notes.\n2. PyMuPDF extracts machine-readable text and clean chapters.\n3. Formats document context for grounded LLM prompts."
    p_sd2.font.name = FONT_MAIN
    p_sd2.font.size = Pt(10)
    p_sd2.font.color.rgb = COLOR_TEXT_BODY
    p_sd2.space_before = Pt(3)

    # Core AI Engine Card
    core_ai = add_card(slide, Inches(13.0), Inches(4.0), Inches(5.6), Inches(1.50), bg_color=COLOR_VERY_LIGHT_CYAN, border_color=COLOR_ACCENT_TEAL, border_width=1.5)
    tf_cai = core_ai.text_frame
    tf_cai.margin_left = Inches(0.2)
    tf_cai.margin_top = Inches(0.15)
    p_cai1 = tf_cai.paragraphs[0]
    p_cai1.text = "🧠 AI Processing Engine (LLM API)"
    p_cai1.font.name = FONT_MAIN
    p_cai1.font.size = Pt(12)
    p_cai1.font.bold = True
    p_cai1.font.color.rgb = COLOR_PRIMARY_NAVY
    p_cai2 = tf_cai.add_paragraph()
    p_cai2.text = "• Receives user doubt + extracted document context\n• Performs grounded semantic reasoning\n• Formats natural language explanations & MCQ questions"
    p_cai2.font.name = FONT_MAIN
    p_cai2.font.size = Pt(10)
    p_cai2.font.color.rgb = COLOR_TEXT_BODY
    p_cai2.space_before = Pt(3)

    # Ingestion to AI Arrow
    arr3 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(12.5), Inches(4.65), Inches(0.4), Inches(0.24))
    arr3.fill.solid()
    arr3.fill.fore_color.rgb = COLOR_ACCENT_TEAL
    arr3.line.fill.background()

    # Dual Outputs Level
    out_qa = add_card(slide, Inches(7.0), Inches(5.95), Inches(5.6), Inches(1.40), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL)
    tf_qa = out_qa.text_frame
    tf_qa.margin_left = Inches(0.2)
    tf_qa.margin_top = Inches(0.15)
    p_qa1 = tf_qa.paragraphs[0]
    p_qa1.text = "💬 Contextual Q&A & Explanation Engine"
    p_qa1.font.name = FONT_MAIN
    p_qa1.font.size = Pt(11.5)
    p_qa1.font.bold = True
    p_qa1.font.color.rgb = COLOR_PRIMARY_NAVY
    p_qa2 = tf_qa.add_paragraph()
    p_qa2.text = "• Answers student doubts grounded in syllabus notes.\n• Breaks down complex formulas into intuitive analogies."
    p_qa2.font.name = FONT_MAIN
    p_qa2.font.size = Pt(10)
    p_qa2.font.color.rgb = COLOR_TEXT_BODY
    p_qa2.space_before = Pt(3)

    out_quiz = add_card(slide, Inches(13.0), Inches(5.95), Inches(5.6), Inches(1.40), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL)
    tf_qz = out_quiz.text_frame
    tf_qz.margin_left = Inches(0.2)
    tf_qz.margin_top = Inches(0.15)
    p_qz1 = tf_qz.paragraphs[0]
    p_qz1.text = "📝 Dynamic Quiz Generator & Scoring"
    p_qz1.font.name = FONT_MAIN
    p_qz1.font.size = Pt(11.5)
    p_qz1.font.bold = True
    p_qz1.font.color.rgb = COLOR_PRIMARY_NAVY
    p_qz2 = tf_qz.add_paragraph()
    p_qz2.text = "• Creates multiple-choice questions from ingested chapters.\n• Delivers immediate evaluation and answer rationales."
    p_qz2.font.name = FONT_MAIN
    p_qz2.font.size = Pt(10)
    p_qz2.font.color.rgb = COLOR_TEXT_BODY
    p_qz2.space_before = Pt(3)

    # Bottom Presentation & Viva Banner
    viva_card = add_card(slide, Inches(7.0), Inches(7.60), Inches(11.6), Inches(2.20), bg_color=COLOR_LIGHT_SURF, border_color=COLOR_CARD_BORDER)
    tf_v = viva_card.text_frame
    tf_v.margin_left = Inches(0.25)
    tf_v.margin_top = Inches(0.18)
    tf_v.word_wrap = True
    p_vh = tf_v.paragraphs[0]
    p_vh.text = "VIVA EXPLAINABILITY & DATA FLOW RATIONALE"
    p_vh.font.name = FONT_MAIN
    p_vh.font.size = Pt(10.5)
    p_vh.font.bold = True
    p_vh.font.color.rgb = COLOR_ACCENT_TEAL
    
    p_vt = tf_v.add_paragraph()
    p_vt.text = "“The student interacts via the React dashboard. Uploaded PDFs are parsed into text using PyMuPDF. When a doubt is submitted, the backend provides the relevant document text alongside the prompt to the LLM API, ensuring syllabus-grounded explanations. Quizzes are generated from the same context for self-assessment.”"
    p_vt.font.name = FONT_MAIN
    p_vt.font.size = Pt(11.5)
    p_vt.font.color.rgb = COLOR_TEXT_DARK
    p_vt.space_before = Pt(4)
    
    p_vs = tf_v.add_paragraph()
    p_vs.text = "Implementation Status: Core PDF ingestion, REST endpoints, and prompt integration under active development (Prototype Architecture)."
    p_vs.font.name = FONT_MAIN
    p_vs.font.size = Pt(10)
    p_vs.font.bold = True
    p_vs.font.color.rgb = COLOR_DEEP_TEAL
    p_vs.space_before = Pt(6)


# =============================================================
# SLIDE 6 BUILDER
# =============================================================
def build_slide_6(slide, assets_dir):
    # Slide 6: Implementation / Prototype
    clear_shapes_except(slide, ['Freeform 2', 'Group 3'])
    add_header(slide, "Implementation / Prototype", "Functional Interface Demonstration & Execution Pipeline")
    add_footer(slide)
    
    # Outer Browser Container
    browser_frame = add_card(slide, Inches(1.0), Inches(1.70), Inches(18.0), Inches(6.55), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    
    # Browser Window Chrome (Top Header Bar)
    b_chrome = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.70), Inches(18.0), Inches(0.48))
    b_chrome.fill.solid()
    b_chrome.fill.fore_color.rgb = COLOR_PRIMARY_NAVY
    b_chrome.line.fill.background()
    
    # Window control dots
    dots = [RGBColor(239, 68, 68), RGBColor(245, 158, 11), RGBColor(16, 185, 129)]
    for d_idx, d_col in enumerate(dots):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2) + d_idx * Inches(0.24), Inches(1.84), Inches(0.14), Inches(0.14))
        dot.fill.solid()
        dot.fill.fore_color.rgb = d_col
        dot.line.fill.background()
        
    # Browser URL Box
    url_box = add_card(slide, Inches(2.2), Inches(1.78), Inches(6.5), Inches(0.32), bg_color=RGBColor(26, 54, 85), border_color=None)
    tf_u = url_box.text_frame
    tf_u.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_u.margin_left = Inches(0.15)
    p_u = tf_u.paragraphs[0]
    p_u.text = "🔒  https://eduverse-ai.local/app/tutor-workspace"
    p_u.font.name = FONT_MAIN
    p_u.font.size = Pt(10)
    p_u.font.color.rgb = COLOR_SOFT_CYAN
    
    # Prototype Interface Pill Badge
    tag_mock = add_card(slide, Inches(14.8), Inches(1.78), Inches(4.0), Inches(0.32), bg_color=COLOR_DEEP_TEAL, border_color=None)
    tf_tm = tag_mock.text_frame
    tf_tm.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_tm = tf_tm.paragraphs[0]
    p_tm.text = "PROTOTYPE INTERFACE (DEMONSTRATION)"
    p_tm.alignment = PP_ALIGN.CENTER
    p_tm.font.name = FONT_MAIN
    p_tm.font.size = Pt(9.5)
    p_tm.font.bold = True
    p_tm.font.color.rgb = COLOR_WHITE
    
    # Browser Body Layout: 3 Columns
    # 1. Left Sidebar (Inches 1.0 to 3.8)
    sidebar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.18), Inches(2.8), Inches(6.07))
    sidebar.fill.solid()
    sidebar.fill.fore_color.rgb = RGBColor(15, 30, 48) # Sleek Dark Navy
    sidebar.line.fill.background()
    
    tf_sb = sidebar.text_frame
    tf_sb.margin_left = Inches(0.22)
    tf_sb.margin_top = Inches(0.24)
    p_sb_brand = tf_sb.paragraphs[0]
    p_sb_brand.text = "EduVerse AI"
    p_sb_brand.font.name = FONT_MAIN
    p_sb_brand.font.size = Pt(16)
    p_sb_brand.font.bold = True
    p_sb_brand.font.color.rgb = COLOR_WHITE
    
    menu_items = [
        ("📊  Dashboard", False),
        ("📚  My Material", True),
        ("🤖  Virtual Tutor", True),
        ("📝  Practice Quizzes", False),
        ("📈  Progress Analytics", False),
        ("⚙️  Settings", False)
    ]
    for mi_text, mi_active in menu_items:
        p_mi = tf_sb.add_paragraph()
        p_mi.text = mi_text
        p_mi.font.name = FONT_MAIN
        p_mi.font.size = Pt(11)
        p_mi.font.bold = mi_active
        p_mi.font.color.rgb = COLOR_SOFT_CYAN if mi_active else RGBColor(180, 195, 210)
        p_mi.space_before = Pt(12)
        
    p_sb_ver = tf_sb.add_paragraph()
    p_sb_ver.text = "\n\nBuild: v0.2-alpha\nStatus: Local Prototype"
    p_sb_ver.font.name = FONT_MAIN
    p_sb_ver.font.size = Pt(9.5)
    p_sb_ver.font.color.rgb = RGBColor(120, 140, 160)

    # 2. Middle Panel: Learning & Dialogue Canvas (Inches 3.9 to 12.4)
    doc_header = add_card(slide, Inches(4.0), Inches(2.35), Inches(8.3), Inches(0.85), bg_color=COLOR_LIGHT_SURF, border_color=COLOR_BORDER_TEAL)
    tf_dh = doc_header.text_frame
    tf_dh.margin_left = Inches(0.2)
    tf_dh.margin_top = Inches(0.12)
    p_dh1 = tf_dh.paragraphs[0]
    p_dh1.text = "📄 ACTIVE DOCUMENT: Machine_Learning_Basics.pdf"
    p_dh1.font.name = FONT_MAIN
    p_dh1.font.size = Pt(11.5)
    p_dh1.font.bold = True
    p_dh1.font.color.rgb = COLOR_PRIMARY_NAVY
    p_dh2 = tf_dh.add_paragraph()
    p_dh2.text = "Unit 4: Supervised Learning • 24 Pages Parsed (PyMuPDF) • Context Grounded"
    p_dh2.font.name = FONT_MAIN
    p_dh2.font.size = Pt(10)
    p_dh2.font.color.rgb = COLOR_DEEP_TEAL
    
    # Dialogue Stream
    # User message bubble
    u_bubble = add_card(slide, Inches(5.2), Inches(3.40), Inches(7.1), Inches(1.10), bg_color=COLOR_VERY_LIGHT_CYAN, border_color=COLOR_CARD_BORDER)
    tf_ub = u_bubble.text_frame
    tf_ub.margin_left = tf_ub.margin_right = Inches(0.2)
    tf_ub.margin_top = Inches(0.12)
    p_ub1 = tf_ub.paragraphs[0]
    p_ub1.text = "👨‍🎓 Student Doubt:"
    p_ub1.font.name = FONT_MAIN
    p_ub1.font.size = Pt(10)
    p_ub1.font.bold = True
    p_ub1.font.color.rgb = COLOR_DEEP_TEAL
    p_ub2 = tf_ub.add_paragraph()
    p_ub2.text = "“Explain supervised learning in simple terms with an example from my notes.”"
    p_ub2.font.name = FONT_MAIN
    p_ub2.font.size = Pt(11.5)
    p_ub2.font.color.rgb = COLOR_PRIMARY_NAVY
    p_ub2.space_before = Pt(2)
    
    # AI response bubble
    ai_bubble = add_card(slide, Inches(4.0), Inches(4.65), Inches(8.3), Inches(2.35), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.2)
    tf_ab = ai_bubble.text_frame
    tf_ab.margin_left = tf_ab.margin_right = Inches(0.22)
    tf_ab.margin_top = Inches(0.14)
    tf_ab.word_wrap = True
    p_ab1 = tf_ab.paragraphs[0]
    p_ab1.text = "🤖 EduVerse AI Tutor (Context-Grounded Explanation):"
    p_ab1.font.name = FONT_MAIN
    p_ab1.font.size = Pt(10)
    p_ab1.font.bold = True
    p_ab1.font.color.rgb = COLOR_ACCENT_TEAL
    
    p_ab2 = tf_ab.add_paragraph()
    p_ab2.text = "“Supervised learning is a machine learning approach where a model learns from labeled training examples (pairs of input features and expected target outputs).\n\n• Analogy: Learning with an answer key.\n• Example from Chapter 4: Predicting real-estate house prices based on known square footage and past transaction records.”"
    p_ab2.font.name = FONT_MAIN
    p_ab2.font.size = Pt(11)
    p_ab2.font.color.rgb = COLOR_TEXT_DARK
    p_ab2.space_before = Pt(3)

    # Action buttons
    btn_ask = add_card(slide, Inches(4.0), Inches(7.20), Inches(3.9), Inches(0.52), bg_color=COLOR_TAG_BG, border_color=COLOR_BORDER_TEAL)
    tf_ba = btn_ask.text_frame
    tf_ba.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_ba = tf_ba.paragraphs[0]
    p_ba.text = "💬  Ask Follow-Up Question"
    p_ba.alignment = PP_ALIGN.CENTER
    p_ba.font.name = FONT_MAIN
    p_ba.font.size = Pt(11)
    p_ba.font.bold = True
    p_ba.font.color.rgb = COLOR_DEEP_TEAL

    btn_quiz = add_card(slide, Inches(8.2), Inches(7.20), Inches(4.1), Inches(0.52), bg_color=COLOR_DEEP_TEAL, border_color=None)
    tf_bq = btn_quiz.text_frame
    tf_bq.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_bq = tf_bq.paragraphs[0]
    p_bq.text = "📝  Generate Concept Practice Quiz"
    p_bq.alignment = PP_ALIGN.CENTER
    p_bq.font.name = FONT_MAIN
    p_bq.font.size = Pt(11)
    p_bq.font.bold = True
    p_bq.font.color.rgb = COLOR_WHITE

    # 3. Right Panel: 3D Avatar & Live Quiz (Inches 12.6 to 18.8)
    avatar_path = os.path.join(assets_dir, 'avatar_assistant.jpg')
    if os.path.exists(avatar_path):
        slide.shapes.add_picture(avatar_path, Inches(12.6), Inches(2.35), Inches(6.15), Inches(3.10))
        
    av_badge = add_card(slide, Inches(12.8), Inches(5.15), Inches(3.6), Inches(0.38), bg_color=COLOR_PRIMARY_NAVY, border_color=None)
    tf_avb = av_badge.text_frame
    tf_avb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_avb = tf_avb.paragraphs[0]
    p_avb.text = "🟢 EduVerse 3D Tutor • Active Session"
    p_avb.font.name = FONT_MAIN
    p_avb.font.size = Pt(9.5)
    p_avb.font.bold = True
    p_avb.font.color.rgb = COLOR_WHITE

    # Quiz card
    quiz_card = add_card(slide, Inches(12.6), Inches(5.65), Inches(6.15), Inches(2.40), bg_color=COLOR_LIGHT_SURF, border_color=COLOR_BORDER_TEAL)
    tf_qc = quiz_card.text_frame
    tf_qc.margin_left = tf_qc.margin_right = Inches(0.2)
    tf_qc.margin_top = Inches(0.14)
    tf_qc.word_wrap = True
    p_qc_h = tf_qc.paragraphs[0]
    p_qc_h.text = "LIVE CONCEPT CHECK (AUTO-GENERATED)"
    p_qc_h.font.name = FONT_MAIN
    p_qc_h.font.size = Pt(10)
    p_qc_h.font.bold = True
    p_qc_h.font.color.rgb = COLOR_ACCENT_TEAL
    
    p_qc_q = tf_qc.add_paragraph()
    p_qc_q.text = "Q1: In supervised learning, what do training samples consist of?"
    p_qc_q.font.name = FONT_MAIN
    p_qc_q.font.size = Pt(11)
    p_qc_q.font.bold = True
    p_qc_q.font.color.rgb = COLOR_PRIMARY_NAVY
    p_qc_q.space_before = Pt(3)
    
    opts = [
        "[   ]  A. Unlabeled data patterns",
        "[ ✓ ]  B. Features paired with ground-truth labels  (Correct)",
        "[   ]  C. Reinforcement reward signals only"
    ]
    for opt in opts:
        p_opt = tf_qc.add_paragraph()
        p_opt.text = opt
        p_opt.font.name = FONT_MAIN
        p_opt.font.size = Pt(10.5)
        p_opt.font.color.rgb = COLOR_SUCCESS if "Correct" in opt else COLOR_TEXT_BODY
        p_opt.space_before = Pt(2)

    # Bottom Implementation Pipeline Banner
    impl_ribbon = add_card(slide, Inches(1.0), Inches(8.40), Inches(18.0), Inches(1.80), bg_color=COLOR_LIGHT_SURF, border_color=COLOR_CARD_BORDER)
    tf_ir = impl_ribbon.text_frame
    tf_ir.margin_left = Inches(0.3)
    tf_ir.margin_top = Inches(0.16)
    p_ir_h = tf_ir.paragraphs[0]
    p_ir_h.text = "IMPLEMENTATION PIPELINE & EXECUTION WORKFLOW"
    p_ir_h.font.name = FONT_MAIN
    p_ir_h.font.size = Pt(10.5)
    p_ir_h.font.bold = True
    p_ir_h.font.color.rgb = COLOR_ACCENT_TEAL
    
    p_ir_sub = tf_ir.add_paragraph()
    p_ir_sub.text = "1. PDF Upload  ➔  2. Text Extraction (PyMuPDF)  ➔  3. AI Context Grounding  ➔  4. Contextual Q&A  ➔  5. Quiz Generation  ➔  6. Instant Evaluation"
    p_ir_sub.font.name = FONT_MAIN
    p_ir_sub.font.size = Pt(12)
    p_ir_sub.font.bold = True
    p_ir_sub.font.color.rgb = COLOR_PRIMARY_NAVY
    p_ir_sub.space_before = Pt(4)
    
    p_ir_note = tf_ir.add_paragraph()
    p_ir_note.text = "Note: Screen demonstrates designed software prototype layout and functional workflow components. Advanced 3D avatar rigging and adaptive analytics are under development."
    p_ir_note.font.name = FONT_MAIN
    p_ir_note.font.size = Pt(10)
    p_ir_note.font.color.rgb = COLOR_TEXT_MUTED
    p_ir_note.space_before = Pt(4)


# =============================================================
# SLIDE 7 BUILDER
# =============================================================
def build_slide_7(slide):
    # Slide 7: Applications & Future Scope
    clear_shapes_except(slide, ['Freeform 2', 'Group 3'])
    add_header(slide, "Applications & Future Scope", "Practical Use Cases & Development Roadmap")
    add_footer(slide)
    
    # Left Column: APPLICATIONS (5 Cards)
    left_card = add_card(slide, Inches(1.0), Inches(1.70), Inches(8.7), Inches(8.5), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf_l = left_card.text_frame
    tf_l.margin_left = tf_l.margin_right = Inches(0.32)
    tf_l.margin_top = Inches(0.24)
    tf_l.word_wrap = True
    
    p_lh = tf_l.paragraphs[0]
    p_lh.text = "PRACTICAL APPLICATIONS"
    p_lh.font.name = FONT_MAIN
    p_lh.font.size = Pt(11)
    p_lh.font.bold = True
    p_lh.font.color.rgb = COLOR_DEEP_TEAL
    
    apps = [
        ("🎓  Student Self-Learning", "Supports independent study sessions, self-paced revision, and instant clarification of complex textbook paragraphs without waiting for office hours."),
        ("📚  Exam Preparation", "Quickly transforms voluminous semester PDFs into structured summaries, key formula explanations, and revision practice sets."),
        ("🔁  Targeted Concept Revision", "Allows students to revisit difficult mathematical or theoretical topics on-demand with simplified breakdowns before laboratory viva."),
        ("📝  Practice Assessment", "Generates instant, topic-wise practice MCQs directly from uploaded chapters to evaluate retention and conceptual understanding."),
        ("💻  Interactive Digital Learning", "Serves as an engaging, multimodal virtual tutor module for institutional learning management systems (LMS).")
    ]
    for a_title, a_desc in apps:
        p_at = tf_l.add_paragraph()
        p_at.text = a_title
        p_at.font.name = FONT_MAIN
        p_at.font.size = Pt(13)
        p_at.font.bold = True
        p_at.font.color.rgb = COLOR_PRIMARY_NAVY
        p_at.space_before = Pt(9)
        
        p_ad = tf_l.add_paragraph()
        p_ad.text = a_desc
        p_ad.font.name = FONT_MAIN
        p_ad.font.size = Pt(11)
        p_ad.font.color.rgb = COLOR_TEXT_BODY
        p_ad.space_before = Pt(2)

    # Right Column: FUTURE SCOPE & ROADMAP
    right_card = add_card(slide, Inches(10.0), Inches(1.70), Inches(9.0), Inches(8.5), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.5)
    tf_r = right_card.text_frame
    tf_r.margin_left = tf_r.margin_right = Inches(0.32)
    tf_r.margin_top = Inches(0.24)
    tf_r.word_wrap = True
    
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "PHASED DEVELOPMENT ROADMAP"
    p_rh.font.name = FONT_MAIN
    p_rh.font.size = Pt(11)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_ACCENT_TEAL
    
    phases = [
        ("PHASE 1: Core AI Tutor (Current Prototype)", "Document parsing via PyMuPDF, document-grounded contextual Q&A, and basic multiple-choice quiz generation."),
        ("PHASE 2: Voice Interaction (Planned)", "Integrating Speech-to-Text (STT) and natural Text-to-Speech (TTS) for natural voice-driven student tutoring."),
        ("PHASE 3: Personalized Learning (Planned)", "Adaptive quiz difficulty that dynamically scales based on individual student accuracy and mastery."),
        ("PHASE 4: Learning Analytics (Planned)", "Long-term student progress tracking, weak-topic identification, and targeted revision recommendations."),
        ("PHASE 5: Cross-Platform Mobile App (Planned)", "Dedicated mobile deployment for smartphone-based portable studying and instant doubt solving.")
    ]
    for p_title, p_desc in phases:
        p_pt = tf_r.add_paragraph()
        p_pt.text = p_title
        p_pt.font.name = FONT_MAIN
        p_pt.font.size = Pt(12)
        p_pt.font.bold = True
        p_pt.font.color.rgb = COLOR_PRIMARY_NAVY
        p_pt.space_before = Pt(9)
        
        p_pd = tf_r.add_paragraph()
        p_pd.text = p_desc
        p_pd.font.name = FONT_MAIN
        p_pd.font.size = Pt(11)
        p_pd.font.color.rgb = COLOR_TEXT_BODY
        p_pd.space_before = Pt(2)

    p_rt_note = tf_r.add_paragraph()
    p_rt_note.text = "\nPlanned Technology Badges:  Speech-to-Text • Text-to-Speech • Multi-Language Support • RAG Document Grounding • Three.js Avatar Lip-Sync"
    p_rt_note.font.name = FONT_MAIN
    p_rt_note.font.size = Pt(10)
    p_rt_note.font.bold = True
    p_rt_note.font.color.rgb = COLOR_DEEP_TEAL
    p_rt_note.space_before = Pt(8)


# =============================================================
# SLIDE 8 BUILDER
# =============================================================
def build_slide_8(slide):
    # Slide 8: Conclusion & References & Thank You
    clear_shapes_except(slide, ['Freeform 2', 'Group 3'])
    add_header(slide, "Conclusion & References", "Project Summary, Documentation & Acknowledgments")
    add_footer(slide)
    
    # Left Column: Conclusion + Takeaway + References (Inches 1.0 to 11.2)
    # Conclusion Card
    conc_card = add_card(slide, Inches(1.0), Inches(1.70), Inches(10.2), Inches(2.65), bg_color=COLOR_WHITE, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf_c = conc_card.text_frame
    tf_c.margin_left = tf_c.margin_right = Inches(0.30)
    tf_c.margin_top = Inches(0.22)
    tf_c.word_wrap = True
    
    p_ch = tf_c.paragraphs[0]
    p_ch.text = "CONCLUSION"
    p_ch.font.name = FONT_MAIN
    p_ch.font.size = Pt(11)
    p_ch.font.bold = True
    p_ch.font.color.rgb = COLOR_DEEP_TEAL
    
    p_ct = tf_c.add_paragraph()
    p_ct.text = "“EduVerse AI demonstrates how generative AI can be integrated with educational content to create an interactive virtual learning assistant. The proposed workflow connects study-material processing, contextual Q&A, explanation generation and quiz-based practice through a unified interface.”"
    p_ct.font.name = FONT_MAIN
    p_ct.font.size = Pt(13)
    p_ct.font.color.rgb = COLOR_TEXT_DARK
    p_ct.space_before = Pt(4)

    # Key Takeaway Banner
    take_card = add_card(slide, Inches(1.0), Inches(4.55), Inches(10.2), Inches(1.30), bg_color=COLOR_PRIMARY_NAVY, border_color=None)
    tf_tk = take_card.text_frame
    tf_tk.margin_left = Inches(0.3)
    tf_tk.margin_top = Inches(0.18)
    p_tkh = tf_tk.paragraphs[0]
    p_tkh.text = "KEY TAKEAWAY WORKFLOW"
    p_tkh.font.name = FONT_MAIN
    p_tkh.font.size = Pt(10)
    p_tkh.font.bold = True
    p_tkh.font.color.rgb = COLOR_SOFT_CYAN
    
    p_tkt = tf_tk.add_paragraph()
    p_tkt.text = "Study Material  ➔  AI Understanding  ➔  Interactive Tutor  ➔  Practice  ➔  Feedback"
    p_tkt.font.name = FONT_MAIN
    p_tkt.font.size = Pt(13.5)
    p_tkt.font.bold = True
    p_tkt.font.color.rgb = COLOR_WHITE
    p_tkt.space_before = Pt(4)

    # References Card
    ref_card = add_card(slide, Inches(1.0), Inches(6.05), Inches(10.2), Inches(4.15), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1.5)
    tf_r = ref_card.text_frame
    tf_r.margin_left = tf_r.margin_right = Inches(0.30)
    tf_r.margin_top = Inches(0.22)
    tf_r.word_wrap = True
    
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "TECHNICAL DOCUMENTATION & REFERENCES"
    p_rh.font.name = FONT_MAIN
    p_rh.font.size = Pt(11)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_ACCENT_TEAL
    
    refs = [
        ("React Documentation", "Official documentation for React component lifecycle & UI architecture (react.dev)."),
        ("FastAPI Documentation", "High-performance Python API framework for asynchronous request handling (fastapi.tiangolo.com)."),
        ("PyMuPDF Documentation", "High-speed PDF text parsing, chapter extraction, and document processing (pymupdf.readthedocs.io)."),
        ("LLM API Documentation", "Large language model interfaces, prompting strategies, and contextual comprehension APIs."),
        ("Three.js Documentation", "WebGL JavaScript 3D library for browser-based character animation & rendering (threejs.org)."),
        ("Academic Literature", "Recent research on AI-assisted personalized tutoring, generative learning interfaces, and automated assessment.")
    ]
    for r_title, r_desc in refs:
        p_rt = tf_r.add_paragraph()
        p_rt.text = f"•  {r_title} — {r_desc}"
        p_rt.font.name = FONT_MAIN
        p_rt.font.size = Pt(10.5)
        p_rt.font.color.rgb = COLOR_TEXT_BODY
        p_rt.space_before = Pt(4)

    # Right Column: OFFICIAL THANK YOU CARD (Inches 11.5 to 19.0)
    ty_card = add_card(slide, Inches(11.5), Inches(1.70), Inches(7.5), Inches(8.5), bg_color=COLOR_LIGHT_SURF, border_color=COLOR_BORDER_TEAL, border_width=1.5)
    tf_ty = ty_card.text_frame
    tf_ty.margin_left = tf_ty.margin_right = Inches(0.40)
    tf_ty.margin_top = Inches(0.40)
    tf_ty.word_wrap = True
    
    p_tyh = tf_ty.paragraphs[0]
    p_tyh.text = "THANK YOU!"
    p_tyh.font.name = FONT_MAIN
    p_tyh.font.size = Pt(36)
    p_tyh.font.bold = True
    p_tyh.font.color.rgb = COLOR_PRIMARY_NAVY
    
    p_ty_proj = tf_ty.add_paragraph()
    p_ty_proj.text = "EduVerse AI"
    p_ty_proj.font.name = FONT_MAIN
    p_ty_proj.font.size = Pt(22)
    p_ty_proj.font.bold = True
    p_ty_proj.font.color.rgb = COLOR_DEEP_TEAL
    p_ty_proj.space_before = Pt(6)
    
    p_ty_sub = tf_ty.add_paragraph()
    p_ty_sub.text = "An Animated Virtual Tutor for Personalized Learning"
    p_ty_sub.font.name = FONT_MAIN
    p_ty_sub.font.size = Pt(13)
    p_ty_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_ty_sub.space_before = Pt(2)

    p_ty_dept = tf_ty.add_paragraph()
    p_ty_dept.text = "\nDepartment of Artificial Intelligence & Data Science\nIndira College of Engineering & Management, Pune\n(An Autonomous Institute Affiliated to SPPU)\nAcademic Year: 2026–2027"
    p_ty_dept.font.name = FONT_MAIN
    p_ty_dept.font.size = Pt(12.5)
    p_ty_dept.font.color.rgb = COLOR_TEXT_DARK
    p_ty_dept.space_before = Pt(12)

    p_ty_tm = tf_ty.add_paragraph()
    p_ty_tm.text = "\n👥  PROJECT TEAM:"
    p_ty_tm.font.name = FONT_MAIN
    p_ty_tm.font.size = Pt(12)
    p_ty_tm.font.bold = True
    p_ty_tm.font.color.rgb = COLOR_ACCENT_TEAL
    p_ty_tm.space_before = Pt(12)
    
    team = ["Aryan Shinde", "Huzaifa Sheikh", "Sufiyan Surve"]
    for tm in team:
        p_tme = tf_ty.add_paragraph()
        p_tme.text = f"•  {tm}"
        p_tme.font.name = FONT_MAIN
        p_tme.font.size = Pt(13)
        p_tme.font.bold = True
        p_tme.font.color.rgb = COLOR_PRIMARY_NAVY
        p_tme.space_before = Pt(3)

    p_ty_mn = tf_ty.add_paragraph()
    p_ty_mn.text = "\n🎓  UNDER THE GUIDANCE OF:"
    p_ty_mn.font.name = FONT_MAIN
    p_ty_mn.font.size = Pt(12)
    p_ty_mn.font.bold = True
    p_ty_mn.font.color.rgb = COLOR_ACCENT_TEAL
    p_ty_mn.space_before = Pt(10)

    p_ty_mp = tf_ty.add_paragraph()
    p_ty_mp.text = "Prof. Monika Patil  (Project Mentor)"
    p_ty_mp.font.name = FONT_MAIN
    p_ty_mp.font.size = Pt(14)
    p_ty_mp.font.bold = True
    p_ty_mp.font.color.rgb = COLOR_PRIMARY_NAVY
    p_ty_mp.space_before = Pt(3)


# =============================================================
# MAIN ORCHESTRATOR
# =============================================================
def main():
    src_path = 's:/SY/SEM 3/MP/EduVerse_AI_Review_2_Presentation.pptx'
    out_path = 's:/SY/SEM 3/MP/EduVerse_AI_Review_2_FINAL.pptx'
    assets_dir = 's:/SY/SEM 3/MP/extracted_assets'
    
    prs = Presentation(src_path)
    print(f"Loaded {src_path} with {len(prs.slides)} slides.")
    
    print("Building Slide 1 (Premium Cover)...")
    build_slide_1(prs.slides[0], assets_dir)
    
    print("Building Slide 2 (Agenda)...")
    build_slide_2(prs.slides[1])
    
    print("Building Slide 3 (Introduction)...")
    build_slide_3(prs.slides[2])
    
    print("Building Slide 4 (Aim & Objectives)...")
    build_slide_4(prs.slides[3])
    
    print("Building Slide 5 (System Architecture & Technology Stack)...")
    build_slide_5(prs.slides[4])
    
    print("Building Slide 6 (Implementation / Prototype)...")
    build_slide_6(prs.slides[5], assets_dir)
    
    print("Building Slide 7 (Applications & Future Scope)...")
    build_slide_7(prs.slides[6])
    
    print("Building Slide 8 (Conclusion & References)...")
    build_slide_8(prs.slides[7])
    
    prs.save(out_path)
    print(f"\nSuccessfully generated and saved refined presentation to:\n{out_path}")

if __name__ == '__main__':
    main()
