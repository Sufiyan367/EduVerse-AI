import pptx
from pptx import Presentation

prs = Presentation('s:/SY/SEM 3/MP/EduVerse_AI_Review_2_Presentation.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"============================== SLIDE {idx+1} ==============================")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    print(f"  [{shape.name}] {txt}")
